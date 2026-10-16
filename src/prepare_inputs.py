"""Prepare authorized local Excel inputs. Never uploads source or derived data."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def write_csv(path, header, rows, delimiter=';'):
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow(header)
        writer.writerows(rows)


def prepare(mkp, mev, output):
    output.mkdir(parents=True, exist_ok=True)
    registry = json.loads((ROOT/'metadata/scope_registry.json').read_text())
    workbook = load_workbook(mkp, read_only=True, data_only=True)
    rows = workbook['бюджет'].iter_rows(values_only=True)
    header = next(rows)
    expected = ['укрупненная группа', 'Регион', 'Вуз', 'Профиль']
    expected += [f'{label} {year}' for label in
                 ['Зачислено на бюджетные места, чел.', 'Зачислено БВИ, чел.',
                  'Ср.балл ЕГЭ (бюджет)'] for year in range(2021, 2026)]
    if list(header[:19]) != expected:
        raise ValueError('Unexpected MKP header/order; inspect the workbook')
    write_csv(output/'mkp_budget.csv', header[:19], (r[:19] for r in rows))
    workbook.close()
    write_csv(output/'scope_links.csv',
              ['university','major_group','scope','evidence_note'],
              ([r['university'],r['major_group'],r['scope'],
                ';'.join(r['source_ids'])] for r in registry['links']), ',')
    inputs = {'mkp': hashlib.sha256(mkp.read_bytes()).hexdigest()}
    if mev:
        mapping = json.loads((ROOT/'metadata/mev_columns.json').read_text())
        workbook = load_workbook(mev, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        columns = [header.index(r['source']) for r in mapping]
        selected = [[r[i] for i in columns] for r in rows]
        write_csv(output/'mev_selected.csv', [r['target'] for r in mapping], selected)
        # Identifiers are obtained from the local source, not published as a dataset.
        identities = json.loads((ROOT/'metadata/university_names.json').read_text())
        matches = []
        for identity in identities:
            candidates = {str(r[0]) for r in selected if r[2] in (None,'')
                          and any(str(r[5]) == name for name in identity['mev_names'])}
            if len(candidates) != 1:
                raise ValueError(f"Ambiguous/missing MEV identity: {identity['short_name']}")
            matches.append([identity['university'],candidates.pop()])
        write_csv(output/'mev_crosswalk.csv',['university','mev_id'],matches,',')
        workbook.close()
        inputs['mev'] = hashlib.sha256(mev.read_bytes()).hexdigest()
    (output/'input_hashes.json').write_text(json.dumps(inputs, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--mkp', type=Path, required=True)
    parser.add_argument('--mev', type=Path)
    parser.add_argument('--output', type=Path, default=ROOT/'data/private')
    args = parser.parse_args()
    prepare(args.mkp, args.mev, args.output)
